"""
gerador.py — gera a AF preenchendo o TEMPLATE OFICIAL (.xlsm) e exporta:
  • Excel: o modelo preenchido.
  • PDF  : via Excel COM (fiel) e **anexa a proposta de origem ao final**
           (AF concluída + proposta). Reserva em reportlab se o Excel faltar.
"""

from __future__ import annotations

import copy
import os
import shutil
import subprocess
import tempfile
import time

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.cell_range import CellRange

from .dados_eletronet import LOGO, MODELO_CPM, template_para
from .log import get_logger
from .modelos import (
    ItemAF, ELETRONET, SEM_MODIFICACAO,
    brl_para_float, float_para_brl, moeda_info, valor_por_extenso,
)

LOG = get_logger("gerador")

ITEM_LIN_INI, ITEM_LIN_FIM = 29, 42
LINHA_VALOR_TOTAL = 43

_CONFIG = {
    "AF": {"originado": "CPM",
           "pg2": {"garantia": "B25", "prazo": "B29", "fat_row": 35,
                   "ent_row": 41, "assinatura": "B53", "nota_precos": "B20"}},
    "AS": {"originado": "CPS",
           "pg2": {"garantia": "B21", "prazo": "B25", "fat_row": 31,
                   "ent_row": 38, "assinatura": "B50", "nota_precos": "B16"}},
}

# reportlab (PDF de reserva, só usado se o Excel COM falhar) é pesado: importado
# SOB DEMANDA dentro de _pdf_reportlab p/ não atrasar a abertura do app.


# ----------------------------------------------------------------- helpers --
def _tipo(prefixo: str) -> str:
    return "AS" if (prefixo or "").upper().startswith("AS") else "AF"


def _parte(prefixo: str) -> str:
    p = (prefixo or "AF-E").split("-")
    return p[1] if len(p) > 1 else "E"


def _sufixo_mod(mod) -> str:
    """"-TR" / "-GE" … ou NADA quando a modificação é "N/A" (ou vem vazia).

    O usuário pode gerar documento sem modificação: aí o identificador fica
    AF-E-444/2026 em vez de AF-E-444/2026-TR."""
    m = str(mod or "").strip()
    return "" if m.upper() in SEM_MODIFICACAO else f"-{m}"


def _af_id(prefixo, numero, ano, mod, revisao="") -> str:
    base = f"{prefixo}-{numero}/{ano}{_sufixo_mod(mod)}"
    rev = str(revisao or "").strip()
    if rev and rev != "0":                       # ex.: AF-E-777/2026-TR-REV111
        base += f"-REV{rev}"
    return base


def _cpm_de(prefixo, numero, ano, mod, revisao="") -> str:
    base = (f"{_CONFIG[_tipo(prefixo)]['originado']}-{_parte(prefixo)}-{numero}/{ano}"
            f"{_sufixo_mod(mod)}")
    rev = str(revisao or "").strip()
    if rev and rev != "0":                       # ex.: CPS-E-789/2026-GE-REV99
        base += f"-REV{rev}"
    return base


def _num(s):
    return brl_para_float(s)


def _itens_ou_padrao(prop) -> list[ItemAF]:
    if prop.itens:
        return prop.itens
    return [ItemAF(codigo="", descricao=(prop.objeto or "Conforme proposta anexada"),
                   quantidade="1", unidade="UN", preco_total_com=prop.valor_total)]


def _remover(path):
    try:
        if path and os.path.exists(path):
            os.remove(path)
    except Exception:
        pass


def _como_lista(valor) -> list[dict]:
    """Normaliza faturamento/entrega: aceita dict único, lista de dicts ou None."""
    if not valor:
        return []
    if isinstance(valor, dict):
        return [valor]
    return [v for v in valor if v]


# ------------------------------------------ página 2: layout dinâmico --------
# O template foi redesenhado (seções movidas, células com fórmulas). Em vez de
# linhas fixas — que ficaram defasadas e gravavam dentro de células mescladas —,
# localizamos cada seção pelo rótulo, como o html_render já faz na leitura.
def _achar_linha(ws, *rotulos) -> int | None:
    alvos = [r.lower() for r in rotulos]
    for row in range(1, ws.max_row + 1):
        for col in range(1, 6):
            v = ws.cell(row, col).value
            if v and any(a in str(v).lower() for a in alvos):
                return row
    return None


def _rows_pg2(ws) -> dict:
    fat = _achar_linha(ws, "DADOS PARA FATURAMENTO")
    ent = _achar_linha(ws, "LOCAIS DE ENTREGA")
    return {
        "garantia": (lambda r: r + 1 if r else None)(_achar_linha(ws, "GARANTIA")),
        "prazo": (lambda r: r + 1 if r else None)(_achar_linha(ws, "PRAZO DE ENTREGA")),
        "fat_data": fat + 2 if fat else None,   # cabeçalho, títulos, dados
        "ent_data": ent + 2 if ent else None,
    }


def _copiar_estilo(src, dst):
    if src.has_style:
        dst.font = copy.copy(src.font)
        dst.border = copy.copy(src.border)
        dst.fill = copy.copy(src.fill)
        dst.number_format = src.number_format
        dst.protection = copy.copy(src.protection)
        dst.alignment = copy.copy(src.alignment)


def _inserir_linhas(ws, data_row: int, n: int):
    """Insere `n` linhas logo abaixo de `data_row`, replicando o estilo/altura da
    linha-modelo e deslocando para baixo as mesclagens e alturas existentes —
    coisas que `openpyxl.insert_rows` sozinho NÃO faz."""
    if n <= 0:
        return
    insert_at = data_row + 1
    movidos = [m for m in list(ws.merged_cells.ranges) if m.min_row >= insert_at]
    for m in movidos:
        ws.merged_cells.ranges.remove(m)
    alturas = {r: d.height for r, d in ws.row_dimensions.items()
               if r >= insert_at and d.height is not None}

    ws.insert_rows(insert_at, n)

    for m in movidos:
        ws.merged_cells.add(str(CellRange(min_col=m.min_col, min_row=m.min_row + n,
                                          max_col=m.max_col, max_row=m.max_row + n)))
    for r in list(alturas):
        ws.row_dimensions[r].height = None
    for r, h in alturas.items():
        ws.row_dimensions[r + n].height = h

    h_modelo = ws.row_dimensions[data_row].height
    colunas = [chr(c) for c in range(ord("A"), ord("U") + 1)]
    for i in range(n):
        r = insert_at + i
        if h_modelo:
            ws.row_dimensions[r].height = h_modelo
        for col in colunas:
            _copiar_estilo(ws[f"{col}{data_row}"], ws[f"{col}{r}"])


def _preencher_secao(ws, data_row: int, locais: list[dict], col_map: list[tuple]):
    """Escreve uma OU várias localidades (faturamento/entrega) empilhadas em
    linhas a partir de `data_row`, inserindo linhas quando há mais de uma."""
    colunas = [c for c, _ in col_map]
    if not locais:
        for col in colunas:                      # remove o exemplo do template
            ws[f"{col}{data_row}"] = None
        return
    _inserir_linhas(ws, data_row, len(locais) - 1)
    for i, loc in enumerate(locais):
        r = data_row + i
        for col, chave in col_map:
            ws[f"{col}{r}"] = loc.get(chave, "") or ""


def _abrir_linhas_de_itens(ws, n_itens: int) -> int:
    """Faz a tabela de itens do modelo CRESCER até caber `n_itens`.

    O modelo tem 14 linhas de item (29-42), e o laço antigo parava na 14ª em
    silêncio: a AF-E-288 da CIENA saiu em Excel com 14 dos 21 itens. Aqui as
    linhas que faltam entram logo abaixo da 42, com o estilo e as mesclagens da
    linha-modelo (C:K, N:O, P:Q — `_inserir_linhas` desloca as de baixo mas não
    cria as novas); o total e as notas descem junto.

    E a folha deixa de ser "caber tudo numa página": o modelo vem assim, e com
    40 itens a letra ficaria ilegível. Passa a ser a LARGURA da folha e a altura
    em quantas folhas precisar, com o cabeçalho da tabela repetido em cada uma.

    Devolve quantas linhas foram acrescentadas (o deslocamento do que vem
    abaixo da tabela)."""
    extra = max(0, n_itens - (ITEM_LIN_FIM - ITEM_LIN_INI + 1))
    if not extra:
        return 0
    _inserir_linhas(ws, ITEM_LIN_FIM, extra)
    for r in range(ITEM_LIN_FIM + 1, ITEM_LIN_FIM + 1 + extra):
        for a, b in (("C", "K"), ("N", "O"), ("P", "Q")):
            ws.merge_cells(f"{a}{r}:{b}{r}")
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0            # 0 = quantas folhas precisar
    ws.print_title_rows = f"{ITEM_LIN_INI - 2}:{ITEM_LIN_INI - 1}"
    return extra


_FAT_COLS = list(zip("BCDEFGH",
                     ("uf", "razao_social", "cnpj", "endereco", "cep", "cnpj2", "insc_est")))
_ENT_COLS = list(zip("BCDEFGHIJ",
                     ("nome", "sigla", "endereco", "municipio", "uf",
                      "maps", "latitude", "longitude", "cedente")))


# ----------------------------------------------------- preenche o template --
def _preencher_template(prop, prefixo, numero, ano, modificacao,
                        data_emissao, revisao, faturamento, entrega, keep_vba=True):
    # keep_vba=True → salva .xlsm (usado no caminho do PDF via Excel COM).
    # keep_vba=False → salva .xlsx VÁLIDO (o Excel rejeita .xlsx com macro embutida).
    faturamentos = _como_lista(faturamento)
    entregas = _como_lista(entrega)
    wb = load_workbook(template_para(prefixo), keep_vba=keep_vba)
    pg1, pg2 = wb["AF-pg1"], wb["AF-pg2"]

    doc_id = _af_id(prefixo, numero, ano, modificacao, revisao)
    cpm = _cpm_de(prefixo, numero, ano, modificacao, revisao)
    # Título (A1): o template AS é cópia do AF e vem com "FORNECIMENTO - AF" — troca p/ AS.
    pg1["A1"] = ("AUTORIZAÇÃO DE SERVIÇO - AS" if _tipo(prefixo) == "AS"
                 else "AUTORIZAÇÃO DE FORNECIMENTO - AF")
    # Cabeçalho da Eletronet nas DUAS folhas. O template trazia na pg2 o endereço
    # da Av. Alfredo Egídio (Chácara Sto. Antônio), de onde a empresa saiu: a pg1
    # dizia Verbo Divino e a pg2 do MESMO documento dizia outro endereço. Escrever
    # a partir de ELETRONET põe o endereço sob controle do código, não do modelo.
    for pg in (pg1, pg2):
        pg["A5"] = ELETRONET["endereco"]
        pg["A6"] = ELETRONET["bairro"]
        pg["A7"] = f'CEP {ELETRONET["cep"]}'

    # ENDEREÇO DE ENTREGA NA PÁGINA 1: o template trazia gravado em O16 o
    # endereço da própria Eletronet em São Paulo, embaixo do rótulo "Endereço de
    # Entrega dos Equipamentos e Materiais". Não era só redundante com a folha 2
    # — era CONTRADITÓRIO: a folha 2 mandava entregar em Fortaleza e a folha 1
    # dizia São Paulo. Nas AS feitas à mão esse campo fica em branco, e o local
    # real vai em "LOCAIS DE ENTREGA" na folha 2. Limpamos, deixando o rótulo.
    pg1["O16"] = None

    valor_f = brl_para_float(prop.valor_total)
    mo = moeda_info(getattr(prop, "moeda", "Real"))
    fmt = mo["fmt"]

    pg1["O10"] = doc_id
    pg1["L11"] = cpm
    pg1["O13"] = data_emissao
    pg1["L17"] = prop.numero_proposta
    pg1["B10"] = prop.fornecedor
    pg1["B12"] = prop.endereco
    pg1["B13"] = prop.cep
    pg1["B14"] = prop.cnpj
    pg1["B15"] = prop.insc_est
    pg1["A19"] = getattr(prop, "moeda", "") or "Real"
    pg1["A22"] = ((valor_por_extenso(valor_f, mo["sing"], mo["plur"], mo["sub"], mo["sub_pl"]) + ".")
                  if valor_f is not None else "")
    pg1["A25"] = prop.objeto

    itens = _itens_ou_padrao(prop)
    extra = _abrir_linhas_de_itens(pg1, len(itens))   # tudo abaixo da tabela desce `extra`
    lin_total = LINHA_VALOR_TOTAL + extra
    for r in range(ITEM_LIN_INI, ITEM_LIN_FIM + extra + 1):
        for col in ("A", "B", "C", "L", "M", "N", "P", "R", "S", "T"):
            pg1[f"{col}{r}"] = None
    from openpyxl.styles import Alignment
    for i, it in enumerate(itens):
        r = ITEM_LIN_INI + i
        pg1[f"A{r}"] = i + 1
        pg1[f"B{r}"] = it.codigo
        pg1[f"C{r}"] = it.descricao
        # DESCRIÇÃO LONGA QUEBRA LINHA. O modelo não quebra: o texto passava
        # da célula e a folha impressa cortava as pontas ("sponder Muxponder
        # ... (lot"). A célula C:K comporta ~110 caracteres em Arial 10
        # negrito; com folga, 95 por linha, e a linha cresce para caber.
        pg1[f"C{r}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        n_linhas = -(-len(str(it.descricao or "")) // 95)
        if n_linhas > 1:
            pg1.row_dimensions[r].height = max(pg1.row_dimensions[r].height or 15, 12.75 * n_linhas + 4)
            # na linha que cresceu, nº, código e preços no MEIO, como a descrição
            # (o modelo os deixa no pé, e o item parecia desalinhado)
            for col in ("A", "B", "L", "M", "N", "P", "R"):
                al = pg1[f"{col}{r}"].alignment
                pg1[f"{col}{r}"].alignment = Alignment(horizontal=al.horizontal, vertical="center",
                                                       wrap_text=al.wrap_text)
        pg1[f"L{r}"] = _num(it.quantidade)
        pg1[f"M{r}"] = it.unidade
        pg1[f"N{r}"] = _num(it.preco_unit_sem)
        pg1[f"P{r}"] = _num(it.preco_unit_com)
        pg1[f"R{r}"] = _num(it.preco_total_com)
        for col in ("N", "P", "R"):
            pg1[f"{col}{r}"].number_format = fmt
    pg1[f"R{lin_total}"] = valor_f
    pg1[f"R{lin_total}"].number_format = fmt
    # U19 (o total do topo) é "=R43" no modelo; o openpyxl não reescreve a
    # fórmula ao inserir linhas, e ela passaria a apontar para um item
    pg1["U19"] = f"=R{lin_total}"
    pg1["U19"].number_format = fmt

    try:
        from datetime import datetime as _dt
        wb["dados_base"]["K2"] = _dt.strptime((data_emissao or "").strip(), "%d/%m/%Y")
    except Exception:
        pass

    # As notas vêm logo abaixo da tabela: descem o que a tabela cresceu.
    def nota(linha):
        return f"B{linha + extra}"

    pg1[nota(45)] = (f"A Proposta da {prop.fornecedor} com a numeração: "
                     f"{prop.numero_proposta}, está anexada e é parte integrante desta.")
    # Observações opcionais (fora do escopo padrão) → linha livre B46 (merge B46:V46).
    obs = [str(o).strip() for o in getattr(prop, "observacoes", []) or [] if str(o).strip()]
    if obs:
        from openpyxl.styles import Alignment
        pg1[nota(46)] = "Observações: " + " • ".join(obs)
        pg1[nota(46)].alignment = Alignment(wrap_text=True, vertical="top")
    pg1[nota(47)] = f"Os preços estão expressos em {mo['plur'].title()} ({mo['simb']})"
    pg1[nota(49)] = "CONDIÇÕES DE FATURAMENTO E PAGAMENTO"
    pg1[nota(50)] = (f"As importâncias objeto desta AF deverão ser pagas pela "
                     f"Eletronet a {prop.fornecedor}, conforme segue:")
    pg1[nota(51)] = prop.condicao_pagamento or ""

    pg2["O10"] = doc_id
    pg2["L11"] = prop.numero_proposta
    pg2["O12"] = data_emissao

    # Seções da pg2 localizadas pelo rótulo (layout do template muda; a
    # assinatura é a fórmula ='AF-pg1'!B10, que já reflete o fornecedor).
    linhas = _rows_pg2(pg2)
    if linhas["garantia"]:
        pg2[f"B{linhas['garantia']}"] = getattr(prop, "garantia", "") or ""
    if linhas["prazo"]:
        pg2[f"B{linhas['prazo']}"] = getattr(prop, "prazo_entrega", "") or ""
    # Preenche ENTREGA (seção de baixo) antes de FATURAMENTO: assim a inserção
    # de linhas no faturamento empurra as linhas de entrega já prontas para
    # baixo sem invalidar índices.
    if linhas["ent_data"]:
        _preencher_secao(pg2, linhas["ent_data"], entregas, _ENT_COLS)
    if linhas["fat_data"]:
        _preencher_secao(pg2, linhas["fat_data"], faturamentos, _FAT_COLS)
    # OBS: a forma de pagamento NÃO é escrita aqui — o template oficial já a
    # traz na pg1 ("CONDIÇÕES DE FATURAMENTO E PAGAMENTO", B49/B51). Escrevê-la
    # também na pg2 fazia aparecer DUAS vezes na AF.

    # Logo: o template tem uma fórmula quebrada (#VALUE!/#REF!) na célula mesclada
    # do canto superior direito — ela aparecia atrás da imagem. Aqui essa célula é
    # LIMPA e o logo entra CENTRADO e enquadrado nela (proporção preservada), em
    # cada aba (pg1: T1:V4, pg2: S1:U4 — colunas diferentes), p/ aparecer só o logo.
    from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.utils.units import pixels_to_EMU
    from openpyxl.utils import get_column_letter

    def _col_px(ws, col):
        return int(round((ws.column_dimensions[get_column_letter(col)].width or 8.43) * 7)) + 5

    def _row_px(ws, row):
        return int(round((ws.row_dimensions[row].height or 15.0) * 4 / 3))

    for ws in (pg1, pg2):
        try:
            merge = next((m for m in ws.merged_cells.ranges
                          if m.min_row == 1 and m.min_col >= 17 and m.max_row >= 3), None)
            if merge is None:
                continue
            ws.cell(merge.min_row, merge.min_col).value = None        # apaga o #VALUE!/#REF!
            mw = sum(_col_px(ws, c) for c in range(merge.min_col, merge.max_col + 1))
            mh = sum(_row_px(ws, r) for r in range(merge.min_row, merge.max_row + 1))
            img = XLImage(LOGO)
            aspecto = (img.width / img.height) if getattr(img, "height", 0) else 1.97
            inset = 6
            cw, ch = mw - 2 * inset, mh - 2 * inset
            if cw / ch > aspecto:                      # sobra largura → limita pela altura
                h = ch; w = h * aspecto
            else:                                      # limita pela largura
                w = cw; h = w / aspecto
            mk = AnchorMarker(col=merge.min_col - 1, colOff=pixels_to_EMU((mw - w) / 2),
                              row=merge.min_row - 1, rowOff=pixels_to_EMU((mh - h) / 2))
            img.anchor = OneCellAnchor(
                _from=mk, ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(h)))
            ws.add_image(img)
        except Exception:
            pass
    return wb


def gerar_af_excel(caminho_saida, prop, prefixo, numero, ano, modificacao,
                   data_emissao, revisao="0", faturamento=None, entrega=None):
    """Gera a AF em .xlsx VÁLIDO. Preferência: Excel COM converte o template .xlsm
    → .xlsx (remove as macros direito). Sem Excel, cai no openpyxl sem VBA."""
    if EXCEL_COM_DISPONIVEL:
        xlsm_tmp = tempfile.NamedTemporaryFile(suffix=".xlsm", delete=False); xlsm_tmp.close()
        try:
            wb = _preencher_template(prop, prefixo, numero, ano, modificacao,
                                     data_emissao, revisao, faturamento, entrega)  # .xlsm p/ o Excel
            wb.save(xlsm_tmp.name)
            if _exportar_xlsx_via_excel(xlsm_tmp.name, caminho_saida):
                return caminho_saida
        finally:
            _remover(xlsm_tmp.name)
    # Reserva (sem Excel COM): openpyxl sem VBA — melhor ter o arquivo que nada.
    wb = _preencher_template(prop, prefixo, numero, ano, modificacao,
                             data_emissao, revisao, faturamento, entrega, keep_vba=False)
    wb.save(caminho_saida)
    return caminho_saida


# -------------------------------------------------- exportação PDF (Excel) --
EXCEL_COM_DISPONIVEL = os.name == "nt"


def _run_powershell(args: list[str], timeout: int = 120) -> tuple[bool, str]:
    """Roda um script PowerShell SEM abrir janela e devolve (ok, erro-resumido).
    Antes as falhas do Excel COM eram engolidas ("except: return False") e o
    fallback acontecia sem pista do motivo — agora o stderr vai para o log."""
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", *args],
                           capture_output=True, text=True, errors="replace",
                           timeout=timeout, creationflags=flags)
    except subprocess.TimeoutExpired:
        return False, f"PowerShell excedeu {timeout}s (Excel travado/aberto com diálogo?)"
    except Exception as exc:
        return False, str(exc)
    err = ((r.stderr or "").strip() or (r.stdout or "").strip())[-800:]
    return r.returncode == 0, err

_PS_EXPORT = '''$ErrorActionPreference = "Stop"
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $false
$xl.DisplayAlerts = $false
try {{ $xl.AutomationSecurity = 3 }} catch {{}}
try {{
  $wb = $xl.Workbooks.Open("{xlsm}", 0, $true)
  $wb.Worksheets("AF-pg1").Select($true)
  $wb.Worksheets("AF-pg2").Select($false)
  $wb.ActiveSheet.ExportAsFixedFormat(0, "{pdf}")
  $wb.Close($false)
}} finally {{
  $xl.Quit()
  [System.Runtime.InteropServices.Marshal]::ReleaseComObject($xl) | Out-Null
}}
'''


def _exportar_pdf_via_excel(xlsm_path, pdf_path) -> bool:
    if not EXCEL_COM_DISPONIVEL:
        return False
    script = _PS_EXPORT.format(xlsm=xlsm_path.replace('"', '`"'), pdf=pdf_path.replace('"', '`"'))
    ps1 = tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8-sig")
    try:
        ps1.write(script)
        ps1.close()
        _remover(pdf_path)
        ok, err = _run_powershell(["-File", ps1.name])
        gerou = os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 0
        if not gerou:
            LOG.error("Excel COM não exportou o PDF (%s): %s", "erro" if not ok else "sem saída",
                      err or "sem detalhes")
        return gerou
    finally:
        _remover(ps1.name)


# Converte o template preenchido (.xlsm) num .xlsx NATIVO via Excel (formato 51 =
# xlOpenXMLWorkbook). O Excel remove as macros corretamente — o openpyxl deixa
# resíduo de "conteúdo habilitado para macro" e o Excel recusa abrir o .xlsx.
_PS_XLSX = '''$ErrorActionPreference = "Stop"
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $false
$xl.DisplayAlerts = $false
try {{ $xl.AutomationSecurity = 3 }} catch {{}}
try {{
  $wb = $xl.Workbooks.Open("{xlsm}", 0, $true)
  # congela as fórmulas das 2 páginas em VALORES (p/ poder remover as abas de dados)
  foreach ($nm in @("AF-pg1","AF-pg2")) {{
    try {{ $ws = $wb.Worksheets($nm); $ws.UsedRange.Value = $ws.UsedRange.Value }} catch {{}}
  }}
  # remove o catálogo interno (dados_base/locais/POPs): enxuga o arquivo e evita
  # vazar 18 fornecedores + 20 filiais + 182 POPs junto com a AF.
  foreach ($nm in @("dados_base","locais","POPs")) {{
    try {{ $wb.Worksheets($nm).Delete() }} catch {{}}
  }}
  $wb.SaveAs("{xlsx}", 51)
  $wb.Close($false)
}} finally {{
  $xl.Quit()
  [System.Runtime.InteropServices.Marshal]::ReleaseComObject($xl) | Out-Null
}}
'''


def _exportar_xlsx_via_excel(xlsm_path, xlsx_path) -> bool:
    if not EXCEL_COM_DISPONIVEL:
        return False
    script = _PS_XLSX.format(xlsm=xlsm_path.replace('"', '`"'), xlsx=xlsx_path.replace('"', '`"'))
    ps1 = tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8-sig")
    try:
        ps1.write(script)
        ps1.close()
        _remover(xlsx_path)
        ok, err = _run_powershell(["-File", ps1.name])
        gerou = os.path.exists(xlsx_path) and os.path.getsize(xlsx_path) > 0
        if not gerou:
            LOG.error("Excel COM não gerou o .xlsx (%s): %s", "erro" if not ok else "sem saída",
                      err or "sem detalhes")
        return gerou
    finally:
        _remover(ps1.name)


def _proposta_pdf(prop) -> str:
    for attr in ("caminho_pdf", "arquivo"):
        caminho = (getattr(prop, attr, "") or "").strip()
        if caminho.lower().endswith(".pdf") and os.path.exists(caminho):
            return caminho
    return ""


def _mesclar_pdfs(af_pdf: str, proposta_pdf: str, saida: str) -> bool:
    try:
        import pypdfium2 as pdfium
    except Exception as exc:
        LOG.warning("pypdfium2 indisponível — proposta não será anexada: %s", exc)
        return False
    dst = af = prop_doc = None
    try:
        dst = pdfium.PdfDocument.new()
        af = pdfium.PdfDocument(af_pdf)
        dst.import_pages(af)
        try:
            prop_doc = pdfium.PdfDocument(proposta_pdf)
            dst.import_pages(prop_doc)
        except Exception:
            LOG.warning("não consegui anexar a proposta %s ao PDF final",
                        os.path.basename(proposta_pdf))
        dst.save(saida)
        return os.path.exists(saida) and os.path.getsize(saida) > 0
    except Exception:
        LOG.exception("falha ao mesclar os PDFs")
        return False
    finally:
        for doc in (prop_doc, af, dst):
            try:
                doc.close()
            except Exception:
                pass


def gerar_af_pdf(caminho_saida, prop, prefixo, numero, ano, modificacao,
                 data_emissao, revisao="0", faturamento=None, entrega=None,
                 anexar_proposta=True, info: dict | None = None):
    """Gera o PDF oficial. `info` (opcional, dict mutável) sai preenchido com o
    motor usado ("excel" fiel ou "reportlab" simplificado) e se a proposta foi
    anexada — o chamador usa isso p/ avisar o usuário sobre fallbacks."""
    if info is None:
        info = {}
    proposta = _proposta_pdf(prop) if anexar_proposta else ""
    info["tinha_proposta"] = bool(proposta)
    xlsm_tmp = tempfile.NamedTemporaryFile(suffix=".xlsm", delete=False); xlsm_tmp.close()
    af_tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False); af_tmp.close()
    gerou = False
    try:
        wb = _preencher_template(prop, prefixo, numero, ano, modificacao,
                                 data_emissao, revisao, faturamento, entrega)
        wb.save(xlsm_tmp.name)
        gerou = _exportar_pdf_via_excel(xlsm_tmp.name, af_tmp.name)
    except Exception:
        LOG.exception("falha ao preencher o template da AF")
        gerou = False
    finally:
        _remover(xlsm_tmp.name)

    info["motor"] = "excel"
    if not gerou:
        LOG.warning("PDF oficial via Excel falhou — caindo no PDF simplificado (reportlab)")
        info["motor"] = "reportlab"
        try:
            _pdf_reportlab(af_tmp.name, prop, prefixo, numero, ano, modificacao, data_emissao, revisao)
            gerou = os.path.exists(af_tmp.name) and os.path.getsize(af_tmp.name) > 0
        except Exception:
            LOG.exception("reportlab também falhou")
            gerou = False
    if not gerou:
        _remover(af_tmp.name)
        return _pdf_reportlab(caminho_saida, prop, prefixo, numero, ano, modificacao, data_emissao, revisao)

    try:
        if proposta and _mesclar_pdfs(af_tmp.name, proposta, caminho_saida):
            info["proposta_anexada"] = True
            return caminho_saida
        info["proposta_anexada"] = False
        shutil.copyfile(af_tmp.name, caminho_saida)
        return caminho_saida
    finally:
        _remover(af_tmp.name)


# -------------------------------------------- "versão HTML": Edge → PDF -------
def _edge_exe() -> str:
    for p in (r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
              r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"):
        if os.path.exists(p):
            return p
    return ""


def gerar_af_pdf_html(caminho_saida, html: str, proposta_pdf: str = "") -> str:
    """Gera o PDF a partir da PRÉVIA HTML da AF (impressa pelo Edge headless) e,
    havendo proposta, anexa-a ao final — espelhando o comportamento do PDF
    oficial. Levanta erro se o Edge não estiver disponível (o chamador faz o
    fallback para o PDF via Excel)."""
    edge = _edge_exe()
    if not edge:
        raise RuntimeError("o Microsoft Edge não foi encontrado nesta máquina")
    html_tmp = tempfile.NamedTemporaryFile("w", suffix=".html", delete=False, encoding="utf-8")
    af_tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False); af_tmp.close()
    gerou = False
    _motivo: list[str] = []
    try:
        html_tmp.write(html); html_tmp.close()
        _remover(af_tmp.name)
        url = "file:///" + html_tmp.name.replace("\\", "/")
        perfil = os.path.join(tempfile.gettempdir(), "geradoraf_edge_pdf")
        # DUAS TENTATIVAS. A 1ª usa o perfil compartilhado (rápido); se falhar,
        # a 2ª usa um perfil descartável, imune a sujeira deixada no primeiro
        # por uma execução anterior que morreu mal. A falha vista na máquina do
        # usuário é intermitente e não se reproduz aqui — contra isso vale
        # insistir, porque a alternativa é o app trocar o LAYOUT do documento
        # sozinho, e quem está padronizando não pode receber ora um, ora outro.
        descartavel = ""
        for tentativa in (1, 2):
            if tentativa == 2:
                descartavel = tempfile.mkdtemp(prefix="geradoraf_edge_")
            usar = descartavel or perfil
            # stdin=DEVNULL não é enfeite: num .exe sem console o processo não
            # tem stdin válido, e o Windows exige as TRÊS entradas válidas
            # assim que uma delas é redirecionada.
            inicio = time.time()
            try:
                r = subprocess.run([edge, "--headless=new", "--disable-gpu",
                                    "--no-pdf-header-footer",
                                    f"--user-data-dir={usar}", "--no-first-run",
                                    f"--print-to-pdf={af_tmp.name}", url],
                                   stdin=subprocess.DEVNULL,
                                   stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                   timeout=120,
                                   creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            except subprocess.TimeoutExpired:
                _motivo.append("tentativa %d: o Edge excedeu 120s" % tentativa)
                r = None
            except OSError as exc:
                _motivo.append("tentativa %d: não consegui executar o Edge (%s)"
                               % (tentativa, exc))
                r = None
            gerou = os.path.exists(af_tmp.name) and os.path.getsize(af_tmp.name) > 0
            if gerou:
                if tentativa == 2:
                    LOG.warning("PDF Visual saiu na 2ª tentativa (perfil limpo): %s",
                                "; ".join(_motivo))
                break
            if r is not None:
                erro = (r.stderr or b"")[-300:].decode("utf-8", "replace").strip()
                _motivo.append("tentativa %d: o Edge terminou em %.0fs com código %s%s"
                               % (tentativa, time.time() - inicio, r.returncode,
                                  (" — " + erro) if erro else ""))
            LOG.error("Edge headless não gerou o PDF: %s", _motivo[-1])
        if descartavel:
            shutil.rmtree(descartavel, ignore_errors=True)
    finally:
        _remover(html_tmp.name)

    if not gerou:
        _remover(af_tmp.name)
        # a razão viaja junto: quem lê o aviso na tela precisa saber o que houve
        raise RuntimeError("; ".join(_motivo) or "o Edge não produziu o PDF")
    try:
        if proposta_pdf and os.path.exists(proposta_pdf) and \
                _mesclar_pdfs(af_tmp.name, proposta_pdf, caminho_saida):
            return caminho_saida
        shutil.copyfile(af_tmp.name, caminho_saida)
        return caminho_saida
    finally:
        _remover(af_tmp.name)


# ------------------------------------------------ reserva: PDF em reportlab --
def _pdf_reportlab(caminho_saida, prop, prefixo, numero, ano, modificacao, data_emissao, revisao="0"):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except Exception:
        raise RuntimeError("reportlab indisponível e Excel COM falhou.")
    mo = moeda_info(getattr(prop, "moeda", "Real"))
    simb = mo["simb"]
    doc = SimpleDocTemplate(caminho_saida, pagesize=landscape(A4),
                            topMargin=12 * mm, bottomMargin=12 * mm,
                            leftMargin=12 * mm, rightMargin=12 * mm)
    azul = colors.HexColor("#1320e0")
    h = ParagraphStyle("h", fontSize=15, leading=18, spaceAfter=6, textColor=azul, fontName="Helvetica-Bold")
    lbl = ParagraphStyle("l", fontSize=8, textColor=colors.grey)
    val = ParagraphStyle("v", fontSize=10, fontName="Helvetica-Bold")
    el = []

    el.append(Paragraph(f"AUTORIZAÇÃO DE {'SERVIÇO - AS' if _tipo(prefixo)=='AS' else 'FORNECIMENTO - AF'}", h))
    el.append(Paragraph(f"Nº {_af_id(prefixo, numero, ano, modificacao, revisao)} &nbsp;·&nbsp; "
                        f"{_cpm_de(prefixo, numero, ano, modificacao, revisao)} &nbsp;·&nbsp; {data_emissao}", lbl))
    el.append(Spacer(1, 8))
    el.append(Paragraph(f"Fornecedor: {prop.fornecedor}", val))
    el.append(Paragraph(f"CNPJ {prop.cnpj} &nbsp; IE {prop.insc_est} &nbsp; {prop.endereco} &nbsp; CEP {prop.cep}", lbl))
    el.append(Spacer(1, 6))
    vf = brl_para_float(prop.valor_total)
    el.append(Paragraph(f"Valor total: {simb} {float_para_brl(vf) if vf is not None else prop.valor_total}", val))
    if vf is not None:
        el.append(Paragraph(valor_por_extenso(vf, mo["sing"], mo["plur"], mo["sub"], mo["sub_pl"]) + ".", lbl))
    el.append(Spacer(1, 6))
    if prop.objeto:
        el.append(Paragraph(f"Objeto: {prop.objeto}", lbl))
        el.append(Spacer(1, 6))

    dados = [["#", "Código", "Descrição", "Qtd", "Un.", f"Unit. ({simb})", f"Total ({simb})"]]
    for i, it in enumerate(_itens_ou_padrao(prop), 1):
        dados.append([str(i), it.codigo, it.descricao, it.quantidade, it.unidade,
                      it.preco_unit_com, it.preco_total_com])
    t = Table(dados, repeatRows=1, colWidths=[18, 70, None, 40, 32, 70, 80])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), azul),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d8deea")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f9fc")]),
    ]))
    el.append(t)
    el.append(Spacer(1, 8))
    el.append(Paragraph(f"Garantia: {prop.garantia or '—'} &nbsp;·&nbsp; "
                        f"Prazo de entrega: {prop.prazo_entrega or '—'}", lbl))
    doc.build(el)
    return caminho_saida


# ============================== CPM / CPS ===================================
# "Coleta de Preços de Material/Serviço" — o relatório técnico-comercial que
# ORIGINA a AF (a AF traz "Originado da Licitação: CPM-...). Gerado nos mesmos
# moldes: preenche o template oficial (assets/modelo_cpm.xlsx) a partir dos
# dados da AF + campos próprios (finalidade, justificativa, consequências…),
# todos editáveis. A alçada (quem aprova) é escolhida pela FAIXA do valor.
_CPM_TETOS = [3000.0, 30000.0, 500000.0, 3000000.0, 10000000.0, float("inf")]


def _alcada(wb, valor: float):
    """Da aba ALÇADA, devolve (rótulo da faixa, [gerente, gerente-geral, diretor,
    cfo, presidência]) conforme o valor — Compras com Concorrência (linhas 11-16).
    O Herivelto saiu da empresa: o cargo de Gerente Geral fica vago (só Leandro)."""
    try:
        al = wb["ALÇADA"]
    except Exception:
        return "", ["", "", "", "", ""]
    idx = next((i for i, teto in enumerate(_CPM_TETOS) if (valor or 0.0) <= teto), 5)
    r = 11 + idx
    rotulo = al.cell(r, 1).value or ""
    nomes = []
    for c in range(2, 7):
        v = al.cell(r, c).value
        v = v if isinstance(v, str) else ""
        if "herivelto" in v.lower():     # cargo de Gerente Geral vago
            v = ""
        nomes.append(v)
    return rotulo, nomes


def _cpm_num(v):
    """Aceita float ou string BR ('20.678,66') e devolve float."""
    if isinstance(v, (int, float)):
        return float(v)
    f = brl_para_float(v)
    return f if f is not None else 0.0


def alcada_para(valor):
    """(rótulo, [gerente, gerente-geral, diretor, cfo, presidência]) p/ o valor —
    lê a aba ALÇADA do template (só leitura). Usado pelo front p/ exibir/editar."""
    from openpyxl import load_workbook
    wb = load_workbook(MODELO_CPM, read_only=True, data_only=True)
    try:
        return _alcada(wb, _cpm_num(valor))
    finally:
        wb.close()


def _ler_faixas(al, linhas, tipo) -> list[dict]:
    out = []
    for r in linhas:
        rotulo = al.cell(r, 1).value or ""
        if not rotulo:
            continue
        nomes = []
        for c in range(2, 7):
            v = al.cell(r, c).value
            v = v if isinstance(v, str) else ""
            if "herivelto" in v.lower():   # Gerente Geral vago (Herivelto saiu)
                v = ""
            nomes.append(v)
        conselho = "Conselho de Administração Eletronet" if "conselho" in rotulo.lower() else ""
        out.append({"label": rotulo, "tipo": tipo, "gerente": nomes[0], "gerente_geral": nomes[1],
                    "diretor": nomes[2], "cfo": nomes[3], "presidencia": nomes[4], "conselho": conselho})
    return out


def alcadas_lista() -> list[dict]:
    """Faixas de alçada das DUAS modalidades (Contratação Direta + Compras com
    Concorrência) p/ a cortina de seleção dos responsáveis."""
    from openpyxl import load_workbook
    wb = load_workbook(MODELO_CPM, read_only=True, data_only=True)
    out = []
    try:
        al = wb["ALÇADA"]
        out += _ler_faixas(al, range(3, 9), "Contratação Direta")       # rows 3-8
        out += _ler_faixas(al, range(11, 17), "Compras com Concorrência")  # rows 11-16
    except Exception:
        pass
    finally:
        wb.close()
    return out


def _cpm_data(s) -> str:
    """Data 'dd/mm/yyyy' → 'yyyy-MM-dd' (p/ o PowerShell montar o datetime)."""
    from datetime import datetime
    try:
        return datetime.strptime(str(s).strip(), "%d/%m/%Y").strftime("%Y-%m-%d")
    except Exception:
        return datetime.today().strftime("%Y-%m-%d")


def _cpm_celulas(cpm: dict):
    """Mapa célula→valor do CPM. Devolve (texto, numero, data). O documento oficial
    é sempre em REAIS: se a moeda for estrangeira, converte pela cotação e a relação
    cambial (H48) explica de onde veio o valor. A alçada é lida da aba ALÇADA."""
    from openpyxl import load_workbook
    eh_real = (moeda_info(cpm.get("moeda", "Real"))["simb"] == "R$")
    taxa = brl_para_float(cpm.get("cotacao", "")) or 0.0

    def _reais(v):   # valor já em R$ (converte se moeda estrangeira)
        n = _cpm_num(v)
        return n if eh_real else n * taxa
    valor = (_cpm_num(cpm.get("valor_reais")) if cpm.get("valor_reais") else _reais(cpm.get("valor")))
    wb = load_workbook(MODELO_CPM, read_only=True, data_only=True)
    try:
        rotulo, nomes = _alcada(wb, valor)
    finally:
        wb.close()
    forn = cpm.get("fornecedor", "") or ""
    forn_full = cpm.get("fornecedor_completo") or forn
    fantasia = cpm.get("fornecedor_fantasia") or forn   # C21 = nome fantasia (DATACOM/FONNET/PADTEC)
    titulo = ("COLETA DE PREÇOS DE SERVIÇO CPS" if cpm.get("tipo") == "CPS"
              else "COLETA DE PREÇOS DE MATERIAL CPM")
    # Aprovadores: editáveis sobrepõem o padrão da alçada (já sem Herivelto).
    aprov = [cpm.get(k) if cpm.get(k) is not None else nomes[i]
             for i, k in enumerate(("gerente", "gerente_geral", "diretor", "cfo", "presidencia"))]
    atend = cpm.get("atendimento") or (
        f"Os equipamentos da {forn_full} atendem as especificações técnicas "
        f"e comerciais solicitadas." if forn_full else "")
    texto = {
        "A1": titulo, "I4": cpm.get("cpm_id", ""), "I13": cpm.get("af_id", ""),
        "A7": cpm.get("finalidade", ""), "A16": cpm.get("justificativa", ""),
        "A30": cpm.get("consequencias", ""), "A21": cpm.get("af_id", ""), "C21": fantasia,
        "A42": atend, "A48": forn_full, "D51": cpm.get("prazo_entrega", ""),
        "D52": cpm.get("local_entrega") or f'{ELETRONET["endereco"]} - {ELETRONET["bairro"]}',
        "D53": cpm.get("forma_pagamento", ""), "D54": cpm.get("alcada") or rotulo,
        "C62": (f"PROPOSTA TÉCNICA E COMERCIAL Proposta: {cpm['proposta']}" if cpm.get("proposta") else ""),
    }
    for cel, nome in zip(("D55", "D56", "D57", "D58", "D59"), aprov):
        texto[cel] = nome or ""
    if not eh_real and cpm.get("relacao_cambial"):   # relação cambial (H48): explica o valor em R$
        texto["H48"] = cpm.get("relacao_cambial")
    capex = _cpm_num(cpm.get("capex", 0))
    numero = {"E21": valor, "H21": valor, "I21": capex,
              "A27": valor, "I27": capex - valor, "I48": valor}   # tudo R$; saldo = CAPEX − total
    # Fornecedores consultados: [nome, prazo, valor]. O template não tem coluna de
    # prazo — vai junto do nome (B); o valor (convertido p/ R$) em I.
    consultados = cpm.get("consultados") or ([[forn_full, "", cpm.get("valor", "")]] if forn_full else [])
    for i in range(4):
        r = 36 + i
        if i < len(consultados):
            row = (list(consultados[i]) + ["", "", ""])[:3] if isinstance(consultados[i], list) else [consultados[i], "", ""]
            nome, prazo, val = row[0], row[1], row[2]
            nome_txt = (nome or "").strip()
            if str(prazo).strip():
                nome_txt = f"{nome_txt} — Prazo: {prazo}".strip(" —")
            texto[f"B{r}"] = nome_txt
            if str(val).strip() != "":
                numero[f"I{r}"] = _reais(val)
            else:
                texto[f"I{r}"] = ""
        else:
            texto[f"B{r}"], texto[f"I{r}"] = "", ""
    data = {"I7": _cpm_data(cpm.get("data_cpm")),
            "I10": _cpm_data(cpm.get("data_relatorio") or cpm.get("data_cpm"))}
    return texto, numero, data


# Preenche o template VIA EXCEL COM (não re-salva pelo openpyxl, que corrompe este
# .xlsx — links externos). Abre o template, escreve as células e exporta PDF/xlsx.
_PS_CPM = r'''param([string]$tmpl,[string]$json,[string]$outpdf,[string]$outxlsx,[string]$logo)
$ErrorActionPreference = "Stop"
$d = Get-Content -Raw -Encoding UTF8 $json | ConvertFrom-Json
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $false
$xl.DisplayAlerts = $false
try { $xl.AutomationSecurity = 3 } catch {}
try {
  $wb = $xl.Workbooks.Open($tmpl, 0, $true)
  $ws = $wb.Worksheets("CPM")
  if ($logo -ne "NONE" -and (Test-Path $logo)) {   # logo BEM ENQUADRADA na célula I1:J3
    try {
      $rm = @(); foreach ($sh in $ws.Shapes) { if ($sh.Type -eq 13) { $rm += $sh } }
      foreach ($sh in $rm) { $sh.Delete() }
      $rng = $ws.Range("I1:J3")
      $pic = $ws.Shapes.AddPicture($logo, $false, $true, $rng.Left, $rng.Top, -1, -1)
      $pic.LockAspectRatio = -1
      $pic.Height = $rng.Height - 8
      if ($pic.Width -gt ($rng.Width - 12)) { $pic.Width = $rng.Width - 12 }
      $pic.Top = $rng.Top + ($rng.Height - $pic.Height) / 2
      $pic.Left = $rng.Left + ($rng.Width - $pic.Width) / 2
    } catch {}
  }
  foreach ($p in $d.texto.PSObject.Properties)  { $ws.Range($p.Name).Value2 = [string]$p.Value }
  foreach ($p in $d.numero.PSObject.Properties) { $ws.Range($p.Name).Value2 = [double]$p.Value }
  foreach ($p in $d.data.PSObject.Properties)   { $ws.Range($p.Name).Value2 = [datetime]::ParseExact([string]$p.Value,'yyyy-MM-dd',$null) }
  if ($d.numfmt -ne "") {   # moeda ≠ Real → troca o símbolo (R$) das células de valor
    foreach ($cel in @("E21","H21","A27","I27","I48","I36","I37","I38","I39")) { try { $ws.Range($cel).NumberFormat = $d.numfmt } catch {} }
  }
  if ($outpdf -ne "NONE")  { $ws.Select($true); $ws.ExportAsFixedFormat(0, $outpdf) }
  if ($outxlsx -ne "NONE") { $wb.SaveAs($outxlsx, 51) }
  $wb.Close($false)
} finally {
  $xl.Quit()
  [System.Runtime.InteropServices.Marshal]::ReleaseComObject($xl) | Out-Null
}
'''


def _gerar_cpm_com(cpm: dict, outpdf="NONE", outxlsx="NONE") -> None:
    import json as _json
    texto, numero, data = _cpm_celulas(cpm)
    numfmt = ""   # documento sempre em R$ (usa o formato R$ do template; sem troca de símbolo)
    jtmp = tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8")
    pstmp = tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8-sig")
    try:
        _json.dump({"texto": texto, "numero": numero, "data": data, "numfmt": numfmt}, jtmp, ensure_ascii=False)
        jtmp.close()
        pstmp.write(_PS_CPM)
        pstmp.close()
        for o in (outpdf, outxlsx):
            if o != "NONE":
                _remover(o)
        logo = os.path.abspath(LOGO) if os.path.exists(LOGO) else "NONE"
        ok, err = _run_powershell(["-File", pstmp.name, os.path.abspath(MODELO_CPM),
                                   jtmp.name, outpdf, outxlsx, logo])
        saiu = any(o != "NONE" and os.path.exists(o) and os.path.getsize(o) > 0
                   for o in (outpdf, outxlsx))
        if not saiu:
            LOG.error("Excel COM não gerou o CPM (%s): %s", "erro" if not ok else "sem saída",
                      err or "sem detalhes")
    finally:
        _remover(jtmp.name)
        _remover(pstmp.name)


def _preencher_cpm_openpyxl(cpm: dict):
    """Reserva (sem Excel COM): preenche via openpyxl. Pode não abrir bem no Excel
    por causa dos links externos do template, mas é o que dá sem o COM."""
    from datetime import datetime
    from openpyxl import load_workbook
    wb = load_workbook(MODELO_CPM)
    try:
        wb._external_links = []
    except Exception:
        pass
    ws = wb["CPM"]
    texto, numero, data = _cpm_celulas(cpm)
    for cel, v in {**texto, **numero}.items():
        ws[cel] = v
    for cel, v in data.items():
        try:
            ws[cel] = datetime.strptime(v, "%Y-%m-%d")
        except Exception:
            pass
    return wb


def gerar_cpm_excel(caminho_saida, cpm: dict):
    if EXCEL_COM_DISPONIVEL:
        _gerar_cpm_com(cpm, outxlsx=caminho_saida)
        if os.path.exists(caminho_saida) and os.path.getsize(caminho_saida) > 0:
            return caminho_saida
        LOG.warning("CPM .xlsx via Excel COM falhou — caindo no openpyxl (pode abrir com aviso)")
    _preencher_cpm_openpyxl(cpm).save(caminho_saida)
    return caminho_saida


def gerar_cpm_pdf(caminho_saida, cpm: dict, proposta_pdf: str = "", info: dict | None = None):
    """Gera o CPM em PDF (Excel COM, fiel ao template) e, havendo proposta, anexa-a
    ao final — igual à AF. Sem o Excel COM, entrega o .xlsx no lugar (o `info`
    reporta motor e anexo p/ o chamador avisar o usuário)."""
    if info is None:
        info = {}
    anexar = bool(proposta_pdf and os.path.exists(proposta_pdf))
    info["tinha_proposta"] = anexar
    info["motor"] = "excel"
    if EXCEL_COM_DISPONIVEL:
        cpm_tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False); cpm_tmp.close()
        _remover(cpm_tmp.name)
        try:
            _gerar_cpm_com(cpm, outpdf=cpm_tmp.name)
            if os.path.exists(cpm_tmp.name) and os.path.getsize(cpm_tmp.name) > 0:
                if anexar and _mesclar_pdfs(cpm_tmp.name, proposta_pdf, caminho_saida):
                    info["proposta_anexada"] = True
                    return caminho_saida                       # CPM + proposta anexada
                info["proposta_anexada"] = False
                shutil.copyfile(cpm_tmp.name, caminho_saida)   # só o CPM (sem/erro no anexo)
                return caminho_saida
        finally:
            _remover(cpm_tmp.name)
    LOG.warning("CPM em PDF indisponível (Excel COM falhou) — entregando .xlsx no lugar")
    info["motor"] = "xlsx"
    saida_xlsx = os.path.splitext(caminho_saida)[0] + ".xlsx"
    _preencher_cpm_openpyxl(cpm).save(saida_xlsx)
    return saida_xlsx
