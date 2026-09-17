"""
extrator_ciena.py — lê a PROPOSTA da CIENA (Excel "DDPTool") e separa em dois:
AF (equipamento/material = CapEx) e AS (serviço = OpEx). A proposta é uma matriz
item × preço-por-estado × quantidade-por-site; usamos só o CÓDIGO, a DESCRIÇÃO,
a QUANTIDADE TOTAL e o TOTAL final de cada bloco (o valor por estado é ignorado,
conforme definido). A CIENA cota em DÓLAR (DDP) — a AF/AS saem em dólar.
"""

from __future__ import annotations

from .modelos import float_para_brl


def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _qtd(v):
    return v if isinstance(v, (int, float)) else None


def _resumo(itens: list[dict]) -> str:
    """Objeto = primeiras palavras das descrições DISTINTAS (igual às AF prontas)."""
    partes: list[str] = []
    for it in itens:
        d = " ".join((it.get("descricao") or "").split()[:5]).strip(" ,")
        if d and d.lower() not in [p.lower() for p in partes]:
            partes.append(d)
    if not partes:
        return ""
    if len(partes) > 1:
        txt = ", ".join(partes[:-1]) + " e " + partes[-1] + "."
    else:
        txt = partes[0] + "."
    return (txt[:297] + "…") if len(txt) > 300 else txt


def ler_ciena(caminho: str) -> dict | None:
    """Devolve {af, as, projeto, moeda, prazo_entrega} se o Excel for uma proposta
    CIENA (abas 'Resumo de Preços' + 'Detalhamento de Preços'); senão, None."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(caminho, data_only=True, read_only=True)
    except Exception:
        return None
    nomes = [s.lower() for s in wb.sheetnames]
    tem_resumo = any("resumo de pre" in n for n in nomes)
    aba_det = next((s for s in wb.sheetnames if "detalhamento de pre" in s.lower()), None)
    if not (tem_resumo and aba_det):
        wb.close()
        return None

    ws = wb[aba_det]
    IC, ID = 2, 3          # C=código, D=descrição (índices 0-based)
    linhas = list(ws.iter_rows(values_only=True))
    # A coluna "Qtde TOTAL" (qtd dos itens E total em US$ das linhas TOTAL) é a
    # ÚLTIMA célula não-vazia da linha de cabeçalho (C="Código"). Varia com o nº de
    # sites (6500=AT, WaveServer5=AK), então achamos dinamicamente.
    IQ = None
    for row in linhas:
        c = _txt(row[IC]) if len(row) > IC else ""
        if c.lower().startswith(("código", "codigo")):
            for j in range(len(row) - 1, ID, -1):
                if _txt(row[j]):
                    IQ = j
                    break
            break
    if IQ is None:
        wb.close()
        return None
    # ---- DESTINOS: uma coluna por site, a UF uma linha acima do nome --------
    # A planilha já traz para onde vai cada coisa. A linha do "Código" tem os
    # NOMES dos sites (da coluna E até a de "Qtde TOTAL"); a linha imediatamente
    # acima tem a UF de cada um.
    entregas: list[dict] = []
    i_cab = next((i for i, row in enumerate(linhas)
                  if len(row) > IC and _txt(row[IC]).lower().startswith(("código", "codigo"))), None)
    if i_cab is not None and i_cab > 0:
        nomes_row, ufs_row = linhas[i_cab], linhas[i_cab - 1]
        for j in range(ID + 1, IQ):
            nome = _txt(nomes_row[j]) if len(nomes_row) > j else ""
            uf = _txt(ufs_row[j]) if len(ufs_row) > j else ""
            if not nome or len(uf) != 2 or not uf.isalpha():
                continue
            # "LICENSES" é a coluna das licenças, não um lugar de entrega
            if nome.strip().upper() in ("LICENSES", "LICENCAS", "LICENÇAS"):
                continue
            entregas.append({"nome": nome.title(), "sigla": "",
                             "municipio": nome.title(), "uf": uf.upper(),
                             "endereco": ""})

    equip: list[dict] = []
    serv: list[dict] = []
    total_equip = total_serv = None
    modo = None            # "equip" | "serv"
    for row in linhas:
        cod = _txt(row[IC]) if len(row) > IC else ""
        desc = _txt(row[ID]) if len(row) > ID else ""
        qtd = _qtd(row[IQ]) if len(row) > IQ else None
        cu, du = cod.upper(), desc.upper()
        if cu.startswith("TOTAL EQUIP"):
            total_equip = qtd if qtd is not None else total_equip
            modo = None
            continue
        if cu.startswith("TOTAL SERV"):
            total_serv = qtd if qtd is not None else total_serv
            modo = None
            continue
        if not cod:                                   # linha de seção/divisória
            if du == "EQUIPAMENTO":
                modo = "equip"
            elif du.startswith("SERVI"):
                modo = "serv"
            continue
        if modo and desc and qtd and qtd > 0:         # item de verdade
            item = {"codigo": cod, "descricao": desc,
                    "quantidade": _txt(qtd), "unidade": "UN",
                    "preco_unit_sem": "", "preco_unit_com": "", "preco_total_com": ""}
            (equip if modo == "equip" else serv).append(item)
    # projeto (aba Resumo, linha do "Projeto ...")
    projeto = ""
    wr = wb[next(s for s in wb.sheetnames if "resumo de pre" in s.lower())]
    for row in wr.iter_rows(values_only=True):
        for v in row:
            t = _txt(v)
            if t.lower().startswith("projeto"):
                projeto = t
                break
        if projeto:
            break
    wb.close()

    if not equip and not serv:
        return None

    def _lado(itens, total):
        return {"itens": itens,
                "valor_total": float_para_brl(total) if total is not None else "",
                "objeto": _resumo(itens)}

    return {
        "ok": True, "eh_ciena": True,
        "moeda": "Dólar Americano",
        "projeto": projeto,
        "prazo_entrega": "24~28 semanas",
        # os destinos valem para os DOIS lados (equipamento e serviço): o
        # serviço é executado onde o equipamento é entregue
        "entregas": entregas,
        "af": _lado(equip, total_equip),   # equipamento/material (CapEx)
        "as": _lado(serv, total_serv),     # serviço (OpEx)
    }
