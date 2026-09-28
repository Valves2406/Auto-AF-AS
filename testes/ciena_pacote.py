# -*- coding: utf-8 -*-
r"""Proposta CIENA: a de um projeto e o PACOTE com vários (cada um, AF + AS).

O QUE MUDOU
-----------
A CIENA passou a mandar um "Bundle": uma aba de pacote e, para CADA projeto, um
par "Resumo X" + "Detalhamento X". O leitor só conhecia "Resumo de Preços" +
"Detalhamento de Preços" e recusava a planilha inteira. Um pacote de 3
projetos são 6 documentos: uma AF (equipamento) e uma AS (serviço) por projeto.

O QUE ESTE TESTE GUARDA
-----------------------
  • o formato de um projeto continua dando 1 AF + 1 AS, como antes;
  • o pacote dá um par por projeto, com o nome e a data de cada um;
  • os itens saem SEM preço (a regra da CIENA: o preço muda por estado), e o
    total de cada documento é o da planilha;
  • quantidade × preço do estado de cada site é CONFERIDO com o total: fecha →
    sem aviso; não fecha → aviso para conferir as quantidades;
  • o lugar do site sai sem o equipamento ("Serra Exemplo WS5" → Serra Exemplo), um por
    lugar; a coluna das licenças não é lugar, mas o estado dela fatura;
  • site sem nome de lugar ("WS5" em SP) não vira entrega — só a UF de faturamento;
  • a tela tem o seletor de projeto na barra da CIENA.

Planilhas montadas aqui, com dados FICTÍCIOS (o repositório é público).
"""
import io
import os
import sys
import tempfile
from datetime import datetime

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
PROJ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(PROJ, "backend"))

import logging
logging.disable(logging.CRITICAL)
from openpyxl import Workbook

from core.extrator_ciena import ler_ciena
from core.modelos import brl_para_float

falhas = []


def ok(nome, cond, extra=""):
    print(("  OK   " if cond else "  FALHA") + " " + nome +
          (("   " + extra) if extra and not cond else ""))
    if not cond:
        falhas.append(nome)


def aba_resumo(wb, nome, titulo, data):
    ws = wb.create_sheet(nome)
    ws["D4"], ws["D5"], ws["D6"], ws["D9"] = "Resumo de Preços", "DATA:", data, titulo
    return ws


def aba_detalhe(wb, nome, titulo, precos, sites, equip, serv, qtd_errada=False):
    """precos: [UF...]; sites: [(nome, UF)...]; itens: (cod, desc, {UF: preço}, [qtd por site])."""
    ws = wb.create_sheet(nome)
    ws["C2"] = titulo
    l_uf, l_cab = 4, 5
    ws.cell(l_cab, 3, "Código"); ws.cell(l_cab, 4, "Descrição")
    col = 5
    cols_preco = {}
    for uf in precos:
        ws.cell(l_cab, col, uf); cols_preco[uf] = col; col += 1
    cols_site = []
    for nome_s, uf in sites:
        ws.cell(l_uf, col, uf); ws.cell(l_cab, col, nome_s); cols_site.append((col, uf)); col += 1
    c_total = col
    ws.cell(l_cab, c_total, "Qtde TOTAL")
    r = l_cab + 1

    def bloco(rotulo, itens, rot_total):
        nonlocal r
        ws.cell(r, 4, rotulo); r += 1
        soma = 0.0
        for k, (cod, desc, pr, qs) in enumerate(itens):
            ws.cell(r, 3, cod); ws.cell(r, 4, desc)
            for uf, p in pr.items():
                ws.cell(r, cols_preco[uf], p)
            for (c, uf), q in zip(cols_site, qs):
                ws.cell(r, c, q)
                soma += q * pr[uf]
            total_q = sum(qs) + (1 if (qtd_errada and k == 0) else 0)   # lida errada de propósito
            ws.cell(r, c_total, total_q)
            r += 1
        ws.cell(r, 3, rot_total); ws.cell(r, c_total, round(soma, 6)); r += 2

    bloco("EQUIPAMENTO", equip, "TOTAL Equipmento (USD) - Com todos os impostos")
    bloco("SERVIÇOS", serv, "TOTAL Serviços (USD) - Com todos os impostos")
    return ws


def item_equip(q_sp_lic, q_sites):
    return [("NTK-EX-001", "AMPLIFICADOR DE LINHA EXEMPLO C-BAND", {"SP": 1000.0, "MG": 1100.0}, q_sites),
            ("LIC-EX-002", "LICENCA DE ESPECTRO EXEMPLO (PER 6.25GHZ)", {"SP": 1.5, "MG": 1.4}, q_sp_lic)]


TMP = tempfile.mkdtemp(prefix="autoaf_ciena_")

# ------------------------------------------------ 1. um projeto só ---------
print("\n1. Planilha de UM projeto (formato de sempre)")
wb = Workbook(); wb.remove(wb.active)
aba_resumo(wb, "Resumo de Preços", "Projeto Exemplo de Ampliação", datetime(2026, 7, 1))
sites = [("Licenses", "SP"), ("Serra Exemplo WS5\n SEX_VEX", "MG"), ("Serra Exemplo 6500\n SEX_VEX", "MG"), ("Granja Exemplo", "SP")]
aba_detalhe(wb, "Detalhamento de Preços", "Projeto Exemplo de Ampliação", ["SP", "MG"], sites,
            equip=[("NTK-EX-001", "AMPLIFICADOR DE LINHA EXEMPLO C-BAND", {"SP": 1000.0, "MG": 1100.0}, [0, 1, 2, 1]),
                   ("LIC-EX-002", "LICENCA DE ESPECTRO EXEMPLO (PER 6.25GHZ)", {"SP": 1.5, "MG": 1.4}, [100, 0, 0, 0])],
            serv=[("SRV-EX-010", "SUPORTE EXEMPLO 2 ANOS", {"SP": 500.0, "MG": 500.0}, [0, 1, 0, 0])])
um = os.path.join(TMP, "um_projeto.xlsx"); wb.save(um)
d = ler_ciena(um)
ok("reconhece como CIENA", bool(d and d.get("eh_ciena")))
if d:
    ok("um projeto", len(d.get("projetos") or []) == 1, str(len(d.get("projetos") or [])))
    ok("AF com os 2 itens de equipamento e AS com o serviço",
       [i["codigo"] for i in d["af"]["itens"]] == ["NTK-EX-001", "LIC-EX-002"]
       and [i["codigo"] for i in d["as"]["itens"]] == ["SRV-EX-010"])
    ok("quantidade total do item", d["af"]["itens"][0]["quantidade"] == "4", d["af"]["itens"][0]["quantidade"])
    ok("total da AF = o da planilha (1000 + 1100×3 + 1,5×100)",
       brl_para_float(d["af"]["valor_total"]) == 4450.0, d["af"]["valor_total"])
    ok("itens SEM preço (o preço muda por estado)",
       all(not i["preco_unit_com"] and not i["preco_total_com"] for i in d["af"]["itens"]))
    ok("a conta fecha e não há aviso", d["af"]["conta_fecha"] and not d["af"]["avisos"], str(d["af"]["avisos"]))
    ok("lugar sem o equipamento, um por lugar; licenças não é lugar",
       [(e["nome"], e["uf"]) for e in d["entregas"]] == [("Serra Exemplo", "MG"), ("Granja Exemplo", "SP")],
       str([(e["nome"], e["uf"]) for e in d["entregas"]]))
    ok("...mas o estado das licenças fatura", d["projetos"][0]["ufs"] == ["SP", "MG"], str(d["projetos"][0]["ufs"]))
    ok("nome e data do projeto", d["projeto"] == "Projeto Exemplo de Ampliação"
       and d["projetos"][0]["data"] == "01/07/2026", "%r %r" % (d["projeto"], d["projetos"][0]["data"]))

# ------------------------------------------------ 2. o pacote --------------
print("\n2. PACOTE com dois projetos (cada um, AF + AS)")
wb = Workbook(); wb.remove(wb.active)
b = wb.create_sheet("Bundle_Exemplo")
b["D5"], b["D6"], b["D9"] = "DATA:", datetime(2026, 9, 1), "Bundle 2 Projetos - Exemplo A / Exemplo B"
aba_resumo(wb, "Resumo Exemplo A", "Projeto Exemplo A", datetime(2026, 8, 2))
aba_resumo(wb, "Resumo Exemplo B", "Transponders Exemplo B", datetime(2026, 8, 3))
aba_detalhe(wb, "Detalhamento Exemplo A", "Projeto Exemplo A", ["SP", "MG"],
            [("Serra Exemplo WS5\n SEX_VEX", "MG"), ("Vale Exemplo 6500\n VEX_SEX", "MG")],
            equip=[("NTK-EX-001", "AMPLIFICADOR DE LINHA EXEMPLO C-BAND", {"SP": 1000.0, "MG": 1100.0}, [1, 1])],
            serv=[("SRV-EX-010", "INSTALACAO EXEMPLO", {"SP": 300.0, "MG": 250.0}, [1, 1])])
aba_detalhe(wb, "Detalhamento Exemplo B", "Transponders Exemplo B", ["SP"],
            [("WS5", "SP"), ("WL6E", "SP")],
            equip=[("WS-EX-100", "WAVESERVER EXEMPLO CHASSIS", {"SP": 2000.0}, [2, 0]),
                   ("WL-EX-200", "WAVESERVER EXEMPLO WL6E", {"SP": 9000.0}, [0, 2])],
            serv=[("SRV-EX-020", "SUPORTE WAVESERVER EXEMPLO", {"SP": 700.0}, [1, 0])])
pac = os.path.join(TMP, "pacote.xlsx"); wb.save(pac)
d = ler_ciena(pac)
ok("reconhece o pacote", bool(d and d.get("eh_ciena")))
if d:
    pj = d.get("projetos") or []
    ok("dois projetos → quatro documentos", len(pj) == 2 and all(p["af"]["itens"] and p["as"]["itens"] for p in pj))
    ok("nome e data de cada projeto",
       [(p["projeto"], p["data"]) for p in pj] == [("Projeto Exemplo A", "02/08/2026"),
                                                  ("Transponders Exemplo B", "03/08/2026")],
       str([(p["projeto"], p["data"]) for p in pj]))
    ok("o pacote: nome e data", d["pacote"].startswith("Bundle 2 Projetos") and d["data_proposta"] == "01/09/2026",
       "%r %r" % (d["pacote"], d["data_proposta"]))
    ok("totais por documento",
       [brl_para_float(p[l]["valor_total"]) for p in pj for l in ("af", "as")] == [2200.0, 500.0, 22000.0, 700.0],
       str([p[l]["valor_total"] for p in pj for l in ("af", "as")]))
    ok("todas as contas fecham", all(p[l]["conta_fecha"] for p in pj for l in ("af", "as")))
    ok("destinos do projeto A", [e["nome"] for e in pj[0]["entregas"]] == ["Serra Exemplo", "Vale Exemplo"])
    ok("projeto B: 'WS5' não é lugar — sem entrega, mas fatura em SP",
       pj[1]["entregas"] == [] and pj[1]["ufs"] == ["SP"] and pj[1]["sem_lugar"], str(pj[1]))
    ok("os campos de sempre trazem o 1º projeto", d["af"] is pj[0]["af"] and d["projeto"] == "Projeto Exemplo A")

# ------------------------------------------------ 3. quantidade lida errada -
print("\n3. Quantidade que não fecha com o total")
wb = Workbook(); wb.remove(wb.active)
aba_resumo(wb, "Resumo de Preços", "Projeto Exemplo C", datetime(2026, 7, 1))
aba_detalhe(wb, "Detalhamento de Preços", "Projeto Exemplo C", ["SP"], [("Granja Exemplo", "SP")],
            equip=[("NTK-EX-001", "AMPLIFICADOR DE LINHA EXEMPLO", {"SP": 1000.0}, [3])],
            serv=[("SRV-EX-010", "SUPORTE EXEMPLO", {"SP": 100.0}, [1])], qtd_errada=True)
err = os.path.join(TMP, "errada.xlsx"); wb.save(err)
d = ler_ciena(err)
ok("não fecha → aviso para conferir as quantidades",
   d and not d["af"]["conta_fecha"] and any("confira as quantidades" in a for a in d["af"]["avisos"]),
   str(d and d["af"]))

# ------------------------------------------------ 4. não é CIENA -------------
wb = Workbook(); wb.active.title = "Planilha1"; wb.active["A1"] = "qualquer coisa"
outra = os.path.join(TMP, "outra.xlsx"); wb.save(outra)
ok("Excel qualquer não vira CIENA", ler_ciena(outra) is None)

# ------------------------------------------------ 5. a tela ----------------
print("\n5. A tela")
js = open(os.path.join(PROJ, "frontend", "app.js"), encoding="utf-8").read()
html = open(os.path.join(PROJ, "frontend", "index.html"), encoding="utf-8").read()
ok("seletor de projeto na barra da CIENA", 'id="cienaProjeto"' in html and "cienaProjeto" in js)
ok("trocar de projeto troca AF/AS daquele projeto", "_ciena.i = +sel.value" in js and "projetoCiena()" in js)
ok("o 🔀 continua trocando AF ↔ AS", 'mostrarLadoCiena(_ciena && _ciena.atual === "af" ? "as" : "af")' in js)

print("FALHAS: " + ", ".join(falhas) if falhas else "TUDO OK")
sys.exit(1 if falhas else 0)
