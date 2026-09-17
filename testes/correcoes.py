# -*- coding: utf-8 -*-
"""Testa (1) o cabecalho das DUAS folhas do Excel e (2) os textos do CPM."""
import sys, io, os, tempfile

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
PROJ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(PROJ, "backend"))
sys.path.insert(0, PROJ)
os.chdir(PROJ)
import engine
from core.modelos import ELETRONET
from openpyxl import load_workbook

falhas = []


def check(nome, obtido, esperado):
    ok = obtido == esperado
    print(("  OK   " if ok else "  FALHA") + " " + nome)
    if not ok:
        print("        obtido:   %r" % (obtido,))
        print("        esperado: %r" % (esperado,))
        falhas.append(nome)


def contem(nome, texto, trecho, deve=True):
    tem = trecho.lower() in (texto or "").lower()
    ok = tem == deve
    print(("  OK   " if ok else "  FALHA") + " " + nome)
    if not ok:
        print("        texto: %r" % ((texto or "")[:150],))
    if not ok:
        falhas.append(nome)


# ============================ 1. cabecalho do Excel ==========================
print("== cabecalho das duas folhas (Excel) ==")
saida = os.path.join(tempfile.gettempdir(), "teste_af_cabecalho.xlsx")
form = {
    "fornecedor": "DATACOM Telemática", "cnpj": "02.820.966/0001-09",
    "endereco": "Rua América, 1000", "cep": "92990-000",
    "valor_total": "10.000,00", "numero": "444", "ano": "2026",
    "prefixo": "AF-E", "modificacao": "N/A", "revisao": "0",
    "objeto": "Fornecimento de switches DmSwitch 2104",
    "itens": [{"codigo": "DM2104", "descricao": "Switch DmSwitch 2104",
               "quantidade": "2", "preco_total_com": "10.000,00"}],
    "formato": "excel", "saida": saida, "nome_arquivo": "teste_af_cabecalho",
}
res = engine.gerar(dict(form))
print("  gerar ->", {k: res[k] for k in ("ok", "erro") if k in res})
arq = res.get("caminho") or saida
if not os.path.exists(arq):
    print("  FALHA: nao gerou o arquivo"); falhas.append("gerar excel")
else:
    wb = load_workbook(arq, data_only=False)
    for aba in ("AF-pg1", "AF-pg2"):
        ws = wb[aba]
        check(f"{aba} A5 endereco", ws["A5"].value, ELETRONET["endereco"])
        check(f"{aba} A6 bairro", ws["A6"].value, ELETRONET["bairro"])
        check(f"{aba} A7 CEP", ws["A7"].value, "CEP " + ELETRONET["cep"])
    print("  pg1 e pg2 iguais:", all(wb["AF-pg1"][c].value == wb["AF-pg2"][c].value
                                     for c in ("A5", "A6", "A7")))

# ============================== 2. textos do CPM =============================
print("\n== textos do CPM ==")
base = {"fornecedor": "DATACOM Telemática", "valor_total": "10.000,00",
        "prefixo": "AF-E", "ano": "2026", "modificacao": "TR", "revisao": "0",
        "cenario": "sem"}

# 2a. o caso da imagem: sem objeto, sem itens, sem numero
d = engine.cpm_de_form(dict(base, numero="", objeto="", itens=[]))
print("  finalidade:", repr(d["finalidade"])[:180])
contem("nao repete 'aquisicao ... realizar a aquisicao'", d["finalidade"],
       "finalidade de realizar a aquisição", deve=False)
contem("nao imprime identificador sem numero", d["finalidade"], "AF-E-/", deve=False)
check("justificativa vazia quando nao ha objeto", d["justificativa"], "")

# 2b. sem objeto, MAS com itens: deduz dos itens
d = engine.cpm_de_form(dict(base, numero="444", objeto="", itens=[
    {"codigo": "DM2104", "descricao": "Switch DmSwitch 2104 com 24 portas"},
    {"codigo": "PSU", "descricao": "Fonte PSU DC 150W"}]))
print("  finalidade:", repr(d["finalidade"])[:200])
contem("usa a descricao dos itens", d["finalidade"], "DmSwitch 2104")
contem("cita o segundo item", d["finalidade"], "PSU DC")
contem("identificador completo entra", d["finalidade"], "AF-E-444/2026-TR")

# 2c. objeto de BEM (lista) -> "para o fornecimento de", nao "realizar 10 switches"
d = engine.cpm_de_form(dict(base, numero="444", objeto="Fornecimento de 10 switches", itens=[]))
print("  finalidade:", repr(d["finalidade"])[:150])
contem("bem usa 'para o fornecimento de'", d["finalidade"], "Aquisição para o fornecimento de 10 switches")
contem("nao escreve 'realizar 10 switches'", d["finalidade"], "realizar 10 switches", deve=False)
check("justificativa de bem", d["justificativa"], "Possibilitar a aquisição de 10 switches.")

# 2d. objeto de ACAO -> mantem "realizar"
d = engine.cpm_de_form(dict(base, numero="444", objeto="a modernização da rede DWDM", itens=[]))
print("  finalidade:", repr(d["finalidade"])[:150])
contem("acao usa 'realizar'", d["finalidade"], "com a finalidade de realizar a modernização da rede DWDM")
check("justificativa de acao", d["justificativa"], "Possibilitar a modernização da rede DWDM.")

# 2e. cenario com cliente segue como era
d = engine.cpm_de_form(dict(base, cenario="com", numero="444", objeto="",
                            clientes_info=[{"cliente": "BMW", "pop_a": "SBA", "pop_b": "MKU",
                                            "banda": "10G", "servico": "transporte"}]))
print("  finalidade:", repr(d["finalidade"])[:190])
contem("cliente preservado", d["finalidade"], "atendimento do cliente BMW")
check("justificativa com cliente", d["justificativa"],
      "Possibilitar o atendimento do cliente BMW nas localidades SBA a MKU.")

print("\n%s" % ("TUDO OK" if not falhas else "FALHAS: " + ", ".join(falhas)))
sys.exit(1 if falhas else 0)
